"""
app.py — AI Job Application Pipeline (Mobile Web App)
Streamlit app using JSearch API (RapidAPI) for job search
and Claude API for fitment analysis, resume rewriting,
networking messages, and interview prep.

Deploy free on Streamlit Cloud:
    1. Push this repo to GitHub
    2. Go to share.streamlit.io → New app → pick your repo
    3. Add ANTHROPIC_KEY and RAPIDAPI_KEY in app Secrets
    4. Open on phone → Add to Home Screen

Local run:
    pip install streamlit anthropic requests openpyxl
    streamlit run app.py
"""

import io
import re
import requests
import streamlit as st
import anthropic
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─── PAGE CONFIG ─────────────────────────────────────────────────────────────

st.set_page_config(
    page_title="AI Job Pipeline",
    page_icon="🤖",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# Mobile-friendly styling
st.markdown("""
<style>
    .block-container { padding-top: 1.5rem; padding-bottom: 3rem; max-width: 720px; }
    .stButton>button { width: 100%; border-radius: 10px; height: 3em; font-weight: 600; }
    div[data-testid="stMetric"] { background: rgba(28,131,225,0.08); border-radius: 10px; padding: 10px; }
    .fit-high  { color: #16a34a; font-weight: 700; }
    .fit-mid   { color: #ca8a04; font-weight: 700; }
    .fit-low   { color: #dc2626; font-weight: 700; }
</style>
""", unsafe_allow_html=True)

# ─── SECRETS / KEYS ──────────────────────────────────────────────────────────

def get_secret(name):
    try:
        return st.secrets[name]
    except Exception:
        import os
        return os.environ.get(name, "")

ANTHROPIC_KEY = get_secret("ANTHROPIC_KEY")
RAPIDAPI_KEY  = get_secret("RAPIDAPI_KEY")

MODEL = "claude-haiku-4-5-20251001"   # cheap + fast; switch to claude-sonnet-4-6 for quality

# ─── SESSION STATE ───────────────────────────────────────────────────────────

if "jobs" not in st.session_state:
    st.session_state.jobs = []
if "analyzed" not in st.session_state:
    st.session_state.analyzed = False

# ─── JSEARCH API ─────────────────────────────────────────────────────────────

def search_jobs(query, location, num_pages=1):
    """Search jobs via JSearch API (RapidAPI)."""
    url = "https://jsearch.p.rapidapi.com/search"
    headers = {
        "X-RapidAPI-Key": RAPIDAPI_KEY,
        "X-RapidAPI-Host": "jsearch.p.rapidapi.com",
    }
    params = {
        "query": f"{query} in {location}",
        "page": "1",
        "num_pages": str(num_pages),
        "country": "in",
        "date_posted": "week",
    }
    try:
        r = requests.get(url, headers=headers, params=params, timeout=30)
        r.raise_for_status()
        data = r.json().get("data", [])
        jobs = []
        for item in data:
            jobs.append({
                "company":  item.get("employer_name") or "Unknown",
                "title":    item.get("job_title") or query,
                "location": f"{item.get('job_city') or location}, {item.get('job_state') or ''}".strip(", "),
                "portal":   item.get("job_publisher") or "JSearch",
                "url":      item.get("job_apply_link") or item.get("job_google_link") or "",
                "jd":       (item.get("job_description") or "")[:3000],
                "posted":   item.get("job_posted_at_datetime_utc", "")[:10],
                "fit": 0, "matches": "", "gaps": "", "verdict": "",
                "resume": "", "cover": "", "outreach": "", "interview": "",
            })
        return jobs, None
    except requests.exceptions.HTTPError as e:
        if e.response.status_code in (401, 403):
            return [], "Invalid RapidAPI key. Check your RAPIDAPI_KEY secret."
        if e.response.status_code == 429:
            return [], "RapidAPI rate limit reached. Free tier allows limited requests/month."
        return [], f"JSearch API error: {e}"
    except Exception as e:
        return [], f"Search failed: {e}"

# ─── CLAUDE AGENTS ───────────────────────────────────────────────────────────

def claude(client, prompt, max_tokens=500):
    msg = client.messages.create(
        model=MODEL,
        max_tokens=max_tokens,
        messages=[{"role": "user", "content": prompt}],
    )
    return msg.content[0].text.strip()


def agent2_fitment(client, job, resume):
    r = claude(client, f"""Analyze resume vs JD. Reply ONLY in this format:
FIT:{{0-100}}
MATCHES:skill1,skill2,skill3
GAPS:gap1,gap2,gap3
VERDICT:one sentence

Resume: {resume}
Company: {job['company']}
JD: {job['jd'][:2000]}""", 300)
    job["fit"]     = int((re.search(r"FIT:(\d+)", r) or [None, 0])[1])
    job["matches"] = (re.search(r"MATCHES:(.+)", r) or [None, ""])[1].strip()
    job["gaps"]    = (re.search(r"GAPS:(.+)",    r) or [None, ""])[1].strip()
    job["verdict"] = (re.search(r"VERDICT:(.+)", r) or [None, ""])[1].strip()


def agent3_resume(client, job, resume, name, title):
    job["resume"] = claude(client, f"""Rewrite resume for this role. Plain text, no markdown.
SUMMARY: [2 sentences tailored to {job['company']}]
BULLETS:
1-4. [achievements with metrics]

Original: {resume}
Target: {title} at {job['company']}
JD: {job['jd'][:1500]}""", 700)


def agent3_cover(client, job, resume, name, title, yrs):
    job["cover"] = claude(client, f"""3-paragraph cover letter. Open with {yrs} yrs experience hook,
middle: 2 achievements matching JD, close with call to action. Sign as {name}.
Resume: {resume}
Company: {job['company']}
JD: {job['jd'][:1000]}""", 450)


def agent5_outreach(client, job, name, title, yrs):
    job["outreach"] = claude(client, f"""Write 2 LinkedIn messages for {job['company']}.
CONNECT: [request to hiring manager, under 280 chars]
FOLLOWUP: [after 7 days no reply, under 250 chars]
Candidate: {name}, {yrs} yrs experience. Role: {title}""", 350)


def agent6_interview(client, job, title, yrs):
    job["interview"] = claude(client, f"""Interview prep for {job['company']}.
T1: [technical question]
T2: [methodology question]
B1: [behavioral - STAR hint]
B2: [stakeholder question]
SD: [case study question]
ASK1: [question to ask them]
ASK2: [question to ask them]
Role: {title}, {yrs} yrs exp. JD: {job['jd'][:1500]}""", 550)

# ─── EXCEL EXPORT ────────────────────────────────────────────────────────────

def build_excel(jobs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fitment Rankings"
    thin, wrap = Side(style="thin", color="CCCCCC"), Alignment(vertical="top", wrap_text=True)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["Rank","Fit%","Company","Title","Location","Portal","Matches","Gaps","Verdict","URL"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = center; cell.border = border

    for i, j in enumerate(jobs, 2):
        fill = PatternFill("solid", fgColor="E2EFDA" if j["fit"]>=85 else "FFEB9C" if j["fit"]>=65 else "FFC7CE")
        for c, v in enumerate([i-1, f"{j['fit']}%", j["company"], j["title"], j["location"],
                               j["portal"], j["matches"], j["gaps"], j["verdict"], j["url"]], 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.fill = fill; cell.border = border
            cell.font = Font(size=9); cell.alignment = center if c in (1,2) else wrap

    for c, w in enumerate([5,6,20,24,16,12,30,28,38,45], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    # Detail sheets for analyzed jobs
    detailed = [j for j in jobs if j.get("resume")]
    if detailed:
        ws2 = wb.create_sheet("Full Prep")
        for c, h in enumerate(["Company","Fit%","Resume","Cover Letter","Outreach","Interview Prep"], 1):
            cell = ws2.cell(row=1, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        for i, j in enumerate(detailed, 2):
            for c, v in enumerate([j["company"], f"{j['fit']}%", j["resume"], j["cover"],
                                   j["outreach"], j["interview"]], 1):
                cell = ws2.cell(row=i, column=c, value=v)
                cell.font = Font(size=9); cell.alignment = wrap
            ws2.row_dimensions[i].height = 140
        for c, w in enumerate([20,6,55,55,45,60], 1):
            ws2.column_dimensions[get_column_letter(c)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

# ─── UI ──────────────────────────────────────────────────────────────────────

st.title("🤖 AI Job Pipeline")
st.caption("Search jobs → AI fitment analysis → tailored resume, outreach & interview prep")

# Key check
missing = []
if not ANTHROPIC_KEY: missing.append("ANTHROPIC_KEY")
if not RAPIDAPI_KEY:  missing.append("RAPIDAPI_KEY")
if missing:
    st.error(f"Missing API keys: {', '.join(missing)}")
    with st.expander("How to fix"):
        st.markdown("""
**Running locally:** set environment variables before launching:
```bash
export ANTHROPIC_KEY=sk-ant-xxxx
export RAPIDAPI_KEY=xxxx
streamlit run app.py
```
**On Streamlit Cloud:** App → Settings → Secrets → add:
```toml
ANTHROPIC_KEY = "sk-ant-xxxx"
RAPIDAPI_KEY = "xxxx"
```
Get a free RapidAPI key: [rapidapi.com/letscrape-6bRBa3QguO5/api/jsearch](https://rapidapi.com/letscrape-6bRBa3QguO5/api/jsearch) → Subscribe to **free Basic plan**.
""")
    st.stop()

# ── Profile form ──
with st.expander("👤 Your Profile", expanded=not st.session_state.jobs):
    col1, col2 = st.columns(2)
    with col1:
        name  = st.text_input("Name", "Nithya")
        title = st.text_input("Job title", "Senior Project Manager")
    with col2:
        yrs      = st.text_input("Years of experience", "20")
        location = st.text_input("Location", "Bengaluru")

    resume = st.text_area("Resume summary", height=140, value=(
        "20 years of experience in end-to-end project management across IT, "
        "infrastructure, and product delivery. PMP-certified. Led cross-functional "
        "teams of 30+, delivered ₹50Cr+ projects on time and within budget. Expert "
        "in Agile, Scrum, Waterfall, JIRA, MS Project, stakeholder management, "
        "vendor negotiations, and C-level reporting across Bengaluru and Coimbatore."
    ))

    max_jobs    = st.slider("Max jobs to fetch", 5, 25, 10)
    top_matches = st.slider("Full prep for top N matches", 1, 5, 3)

# ── Step 1: Search ──
if st.button("🔍 Step 1 — Search Jobs", type="primary"):
    with st.spinner(f"Searching {title} jobs in {location}..."):
        jobs, err = search_jobs(title, location)
        if err:
            st.error(err)
        elif not jobs:
            st.warning("No jobs found. Try a broader title or different location.")
        else:
            st.session_state.jobs = jobs[:max_jobs]
            st.session_state.analyzed = False
            st.success(f"Found {len(st.session_state.jobs)} jobs!")

# ── Show search results ──
if st.session_state.jobs and not st.session_state.analyzed:
    st.subheader(f"📋 {len(st.session_state.jobs)} Jobs Found")
    for j in st.session_state.jobs:
        st.markdown(f"**{j['company']}** — {j['title']}  \n"
                    f"📍 {j['location']} · via {j['portal']} · {j['posted']}")
    st.divider()

    # ── Step 2: Analyze ──
    if st.button("🎯 Step 2 — Run AI Fitment on All Jobs", type="primary"):
        client = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
        prog = st.progress(0, "Starting analysis...")
        errors = 0
        for idx, job in enumerate(st.session_state.jobs):
            prog.progress((idx+1)/len(st.session_state.jobs),
                          f"Analyzing {job['company']} ({idx+1}/{len(st.session_state.jobs)})")
            try:
                agent2_fitment(client, job, resume)
            except Exception as e:
                job["verdict"] = f"Error: {e}"
                errors += 1
        st.session_state.jobs.sort(key=lambda x: x["fit"], reverse=True)
        st.session_state.analyzed = True
        prog.empty()
        if errors:
            st.warning(f"{errors} jobs failed analysis (check API key/credits)")
        st.rerun()

# ── Show analyzed results ──
if st.session_state.analyzed and st.session_state.jobs:
    jobs = st.session_state.jobs
    strong = sum(1 for j in jobs if j["fit"] >= 85)
    good   = sum(1 for j in jobs if 65 <= j["fit"] < 85)

    st.subheader("🏆 Fitment Rankings")
    c1, c2, c3 = st.columns(3)
    c1.metric("Analyzed", len(jobs))
    c2.metric("Strong ≥85%", strong)
    c3.metric("Good 65-84%", good)

    for i, j in enumerate(jobs, 1):
        cls = "fit-high" if j["fit"] >= 85 else "fit-mid" if j["fit"] >= 65 else "fit-low"
        with st.expander(f"#{i}  {j['company']} — {j['fit']}% fit"):
            st.markdown(f"**{j['title']}** · {j['location']} · via {j['portal']}")
            st.markdown(f"<span class='{cls}'>Fit: {j['fit']}%</span>", unsafe_allow_html=True)
            if j["matches"]: st.markdown(f"✅ **Matches:** {j['matches']}")
            if j["gaps"]:    st.markdown(f"⚠️ **Gaps:** {j['gaps']}")
            if j["verdict"]: st.markdown(f"💬 {j['verdict']}")
            if j["url"]:     st.link_button("Apply →", j["url"])

            # Show full prep if generated
            if j.get("resume"):
                st.divider()
                t1, t2, t3, t4 = st.tabs(["📝 Resume", "✉️ Cover", "🤝 Outreach", "🎤 Interview"])
                with t1: st.text(j["resume"])
                with t2: st.text(j["cover"])
                with t3: st.text(j["outreach"])
                with t4: st.text(j["interview"])

    st.divider()

    # ── Step 3: Full prep on top N ──
    top = [j for j in jobs if j["fit"] >= 60][:top_matches]
    if top and not top[0].get("resume"):
        if st.button(f"🚀 Step 3 — Full Prep for Top {len(top)} Companies", type="primary"):
            client = anthropic.Anthropic(api_key=ANTHROPIC_KEY)
            prog = st.progress(0, "Generating...")
            total = len(top) * 4
            step = 0
            for job in top:
                for fn, label in [
                    (lambda: agent3_resume(client, job, resume, name, title), "resume"),
                    (lambda: agent3_cover(client, job, resume, name, title, yrs), "cover letter"),
                    (lambda: agent5_outreach(client, job, name, title, yrs), "outreach"),
                    (lambda: agent6_interview(client, job, title, yrs), "interview prep"),
                ]:
                    step += 1
                    prog.progress(step/total, f"{job['company']}: {label}...")
                    try:
                        fn()
                    except Exception as e:
                        st.warning(f"{job['company']} {label} failed: {e}")
            prog.empty()
            st.success("Full prep complete! Expand each company above to see everything.")
            st.rerun()

    # ── Download Excel ──
    st.download_button(
        "📊 Download Full Excel Report",
        data=build_excel(jobs),
        file_name=f"job_pipeline_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    if st.button("🔄 New Search"):
        st.session_state.jobs = []
        st.session_state.analyzed = False
        st.rerun()
